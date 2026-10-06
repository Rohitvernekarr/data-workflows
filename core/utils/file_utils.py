"""Resolve and inspect input workbooks."""

from __future__ import annotations

import io
import os
import posixpath
import re
from pathlib import Path
from typing import Optional
import configparser
import pandas as pd

import openpyxl
import paramiko

CONFIG_PATH = Path(__file__).resolve().parents[1] / "config.ini"

config = configparser.ConfigParser()
config.read(CONFIG_PATH)

BASE_INPUT_FOLDER = config.get("sftp", "input_files_location", fallback="/home/nifi/nifi-files/verisure/input_files")
MAPPING_REVIEW_FOLDER = config.get("sftp", "mapping_review_files_location", fallback="/home/nifi/nifi-files/verisure/mapping_review_files")
MAPPED_FILES_FOLDER = config.get("sftp", "mapped_files_location", fallback="/home/nifi/nifi-files/verisure/mapped_files")
TRANSFORMED_FILES_FOLDER = config.get("sftp", "transformed_files_location", fallback="/home/nifi/nifi-files/verisure/transformed_files")
TRANSFORMED_FILES_BACKUP_FOLDER = config.get("sftp", "transformed_files_backup_location", fallback="/home/nifi/nifi-files/verisure/transformed_files_backup")


def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (s or "").lower())


def _looks_like_path(value: str) -> bool:
    if not value:
        return False
    # Treat Windows, Unix, and relative-style inputs as already-qualified paths.
    return bool(
        re.match(r"^[a-zA-Z]:[\\/]", value)
        or value.startswith(("/", "\\", "~"))
        or "/" in value
        or "\\" in value
    )


def _is_sftp_path(value: str) -> bool:
    value = (value or "").strip()
    # Relative SFTP locations should be treated as remote paths, not local files.
    return (
        bool(value)
        and value.startswith(("/", "./", "../"))
        and not re.match(r"^[a-zA-Z]:[\\/]", value)
    )


def _resolve_file_reference(
    filename: str, source_files_location: str, sftp_remote: bool
) -> str:
    filename = (filename or "").strip()
    source_files_location = (source_files_location or "").strip()

    if not filename:
        return ""

    if _looks_like_path(filename):
        return filename

    if not source_files_location:
        return filename

    # Bare filenames are attached to the configured source location.
    # Explicit paths are preserved as-is for both local and SFTP inputs.
    # This keeps the caller's local path or remote path semantics intact.
    if sftp_remote or _is_sftp_path(source_files_location):
        return posixpath.join(source_files_location.rstrip("/"), filename)

    return str(Path(source_files_location) / filename)


def _read_file_bytes(
    file_reference: str, sftp: Optional[paramiko.SFTPClient] = None
) -> bytes:
    if os.path.exists(file_reference):
        return Path(file_reference).read_bytes()

    if sftp is None:
        raise FileNotFoundError(f"Local file not found: {file_reference}")

    with io.BytesIO() as buf:
        sftp.getfo(file_reference, buf)
        return buf.getvalue()



def _sheet_names(file_bytes: bytes) -> list[str]:
    """Read workbook tabs without changing the input file."""
    try:
        wb = openpyxl.load_workbook(
            io.BytesIO(file_bytes), read_only=True, data_only=True
        )
        try:
            return wb.sheetnames
        finally:
            wb.close()
    except Exception:
        return []


def _extract_file_info(
    file_bytes: bytes,
    skip_rows: int = 0,
    target_sheet: str = "",
) -> tuple[list[str], list[str]]:
    """
    Returns (sheet_names, headers).

    - sheet_names : all tab names in the workbook
    - headers     : column headers read from row (skip_rows + 1) of the
                    target_sheet (or active sheet if not found / not given)

    File is opened read-only and never modified.
    """
    sheet_names: list[str] = []
    headers: list[str] = []

    try:
        wb = openpyxl.load_workbook(
            io.BytesIO(file_bytes), read_only=True, data_only=True, keep_links=False
        )
        sheet_names = wb.sheetnames

        ws = None
        if target_sheet:
            if target_sheet in wb.sheetnames:
                ws = wb[target_sheet]
                # get_run_logger().debug("Using configured sheet (exact): '%s'", target_sheet)
            else:
                for sn in wb.sheetnames:
                    if _norm(sn) == _norm(target_sheet):
                        ws = wb[sn]
                        # get_run_logger().debug("Using configured sheet (case-insensitive): '%s'", sn)
                        break

        if ws is None:
            ws = wb.active
            # get_run_logger().debug(
            #     "Sheet '%s' not found or not configured — using active sheet: '%s'",
            #     target_sheet, ws.title,
            # )

        header_row = skip_rows + 1  # openpyxl rows are 1-based
        headers = [
            str(ws.cell(row=header_row, column=c).value).strip()
            for c in range(1, ws.max_column + 1)
            if ws.cell(row=header_row, column=c).value not in (None, "")
        ]
        wb.close()

    except Exception as e:
        pass
        # get_run_logger().warning("File inspection failed: %s", e)

    # get_run_logger().debug("FILE — Sheets      : %s", sheet_names)
    # get_run_logger().debug("FILE — skip_rows=%d  header_row=%d  Headers: %s",
    # skip_rows, skip_rows + 1, headers)
    return sheet_names, headers


def build_input_file_path(partner_name: str) -> str:
    """
    Return a valid Linux-style path under a fixed folder.
    The input string is sanitized and appended as a suffix.
    """
    safe_name = get_safe_name(partner_name)
    return str(Path(BASE_INPUT_FOLDER) / safe_name)

def build_mapped_file_path(partner_name: str) -> str:
    """
    Return a valid Linux-style path under a fixed folder.
    The output string is sanitized and appended as a suffix.
    """
    safe_name = get_safe_name(partner_name)
    return str(Path(MAPPED_FILES_FOLDER) / safe_name)

def build_mapping_review_file_path(partner_name: str) -> str:
    """
    Return a valid Linux-style path under a fixed folder.
    The output string is sanitized and appended as a suffix.
    """
    safe_name = get_safe_name(partner_name)
    return str(Path(MAPPING_REVIEW_FOLDER) / safe_name)

def build_transformed_file_path(partner_name: str) -> str:
    """
    Return a valid Linux-style path under a fixed folder.
    The output string is sanitized and appended as a suffix.
    """
    safe_name = get_safe_name(partner_name)
    return str(Path(TRANSFORMED_FILES_FOLDER) / safe_name)

def build_transformed_files_backup_file_path(partner_name: str) -> str:
    """
    Return a valid Linux-style path under a fixed folder.
    The output string is sanitized and appended as a suffix.
    """
    safe_name = get_safe_name(partner_name)
    return str(Path(TRANSFORMED_FILES_BACKUP_FOLDER) / safe_name)

def get_safe_name(partner_name: str) -> str:
    """
    Return a sanitized version of the partner name.
    """
    return "".join(
        ch if ch not in { "/", "\\" } else "-"
        for ch in partner_name.strip()
    ).strip("-")

class FileUtils:
    """Stateless file access and workbook inspection helpers."""

    resolve_file_reference = staticmethod(_resolve_file_reference)
    read_file_bytes = staticmethod(_read_file_bytes)
    sheet_names = staticmethod(_sheet_names)
    extract_file_info = staticmethod(_extract_file_info)
